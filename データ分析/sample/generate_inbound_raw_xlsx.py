"""訪日外客数のデータ加工用に、JNTO「時系列推移表」のレイアウトを模した加工前Excel(架空値)を生成する。

実行: python generate_inbound_raw_xlsx.py  (openpyxl が必要)
- 1シート=1年分。上位3行はタイトル・空行・注記。4行目が見出し行(月と伸率が交互、末尾に累計と伸率)
- 2019年以前: 国名はA列のみ
- 2020年以降: 国名がA列またはB列(B列は先頭に全角スペース付きの下位項目)
- 最新年(2023)は5月分まで入力済みで以降は空欄
- 国名の前後スペース、地域計・総数・注の行、伸率・累計の列を含む
数値はすべて乱数による架空値で、実際の統計値ではない。
"""
import random
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font

OUT = Path(__file__).parent / "data"
OUT.mkdir(exist_ok=True)

LEAVES = {  # 国名: 月平均の基準人数
    "韓国": 480000, "中国": 420000, "台湾": 400000, "香港": 170000, "タイ": 90000, "シンガポール": 32000,
    "イスラエル": 3000, "トルコ": 2000, "その他アジア": 60000, "英国": 30000, "フランス": 29000, "ドイツ": 24000,
    "その他ヨーロッパ": 60000, "米国": 130000, "カナダ": 28000, "豪州": 50000,
}
GROUPS = [
    ("アジア計", ["韓国", "中国", "台湾", "香港", "タイ", "シンガポール", "イスラエル", "トルコ", "その他アジア"]),
    ("ヨーロッパ計", ["英国", "フランス", "ドイツ", "その他ヨーロッパ"]),
    ("北アメリカ計", ["米国", "カナダ"]),
    ("オセアニア計", ["豪州"]),
]
SEASON = [0.85, 0.9, 1.1, 1.15, 1.0, 0.95, 1.05, 1.1, 0.95, 1.1, 1.05, 1.0]
LATEST_YEAR, LATEST_MONTHS = 2023, 5


def factor(y, m):
    if y == 2020:
        return 0.95 if m <= 2 else 0.3 if m == 3 else 0.01
    if y == 2021:
        return 0.002
    if y == 2022:
        return 0.04 + 0.03 * m / 12 if m < 6 else 0.1 + 0.03 * (m - 6)
    if y == 2023:
        return 0.45 + 0.1 * m if m <= 7 else 0.9
    return {2017: 0.85, 2018: 0.95, 2019: 1.0}[y]


def main():
    r = random.Random(707)
    data = {(y, c): [int(b * SEASON[m - 1] * factor(y, m) * r.gauss(1, 0.04)) for m in range(1, 13)]
            for y in range(2017, 2024) for c, b in LEAVES.items()}

    def total(y, names):
        return [sum(data[(y, n)][m] for n in names) for m in range(12)]

    def series(y, label):
        if label == "総数":
            return total(y, list(LEAVES))
        if label == "中東地域":
            return total(y, ["イスラエル", "トルコ"])
        for g, names in GROUPS:
            if label == g:
                return total(y, names)
        return data[(y, label)]

    def growth(cur, prev):
        return round((cur - prev) / prev * 100, 1) if prev else None

    wb = Workbook()
    wb.remove(wb.active)
    for y in (2023, 2022, 2020, 2019, 2018):
        ws = wb.create_sheet(str(y))
        new_layout = y >= 2020
        pad = 2 if new_layout else 1
        ws.append([f"{y}年\u3000訪日外客数（総数）"])
        ws.append([])
        ws.append(["＊本表で、通年の月別・市場別の推移が確認できます。伸率は前年同月比を表しています。"])
        head = [None] * pad
        for m in range(1, 13):
            head += [f"{m}月", "伸率"]
        ws.append(head + ["累計", "伸率"])
        order = ["総数", "アジア計", "韓国", "中国", "台湾", "香港", "タイ", "シンガポール"]
        order += ["中東地域", "\u3000イスラエル", "\u3000トルコ"] if new_layout else ["イスラエル", "トルコ"]
        order += ["その他アジア", "ヨーロッパ計", "英国", "フランス", "ドイツ", "その他ヨーロッパ", "北アメリカ計",
                  "米国 ", "カナダ", "オセアニア計", "豪州"]
        filled = LATEST_MONTHS if y == LATEST_YEAR else 12
        for label in order:
            cur, prev = series(y, label.strip()), series(y - 1, label.strip())
            row = [None] * pad
            if label.startswith("\u3000"):
                row[1] = label
            else:
                row[0] = label
            for m in range(12):
                row += [cur[m], growth(cur[m], prev[m])] if m < filled else [None, None]
            row += [sum(cur[:filled]), growth(sum(cur[:filled]), sum(prev[:filled]))]
            ws.append(row)
            if y == LATEST_YEAR:
                ws.cell(ws.max_row, pad + 2 * (LATEST_MONTHS - 1) + 1).font = Font(italic=True)
        ws.append(["注１：\u3000本資料は研修用の架空データです。実際の統計値ではありません。"])
        ws.append(["注２：\u3000" + ("斜体の数値は推計値である。" if y == LATEST_YEAR else "表中の数値は、全て確定値である。")])
    out = OUT / "inbound_visitors_raw_layout_sample.xlsx"
    wb.save(out)
    print(f"{out.name}: sheets={wb.sheetnames}")


if __name__ == "__main__":
    main()
